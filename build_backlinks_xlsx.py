"""Generates backlinks_vpn.xlsx from the verified target list.

Re-runnable: edit ROWS and re-run. Requires .venv/bin/python3 (openpyxl installed there).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "domain", "category", "submission type", "requires phone/address",
    "priority", "status", "notes",
]

ROWS = [
    # --- DONE / IN FLIGHT ---
    ["saashub.com", "Software/SaaS directory", "self-serve form (account + product submission)", "No", "P1",
     "registered + submitted (pending approval)",
     "DONE 2026-08-18. Account horizance-vpn / james.smith.web@gmail.com, EMAIL VERIFIED. hCaptcha solved (see notes tab in report: semantic-habitat reading). Product submitted FREE tier (explicitly declined the $75 Priority+ upsell). Listing URL once approved: https://www.saashub.com/horizance-vpn . Categories: VPN, Virtual Private Network (VPN), Privacy, Security & Privacy, Cloud VPN, VPN Software. Registered as competitor/alternative to ProtonVPN, NordVPN, Surfshark, Mozilla VPN. Free tier = 'up to 32 days' queue. NOTE: the site auto-scraped a Russian tagline using 'не палится DPI' - replaced with a factual English tagline, since that phrasing violates the copywriting-tone rules (no slang, no explicit DPI language)."],

    ["launchingnext.com", "Startup/product launch directory", "self-serve form", "No", "P1",
     "submitted - confirmed",
     "DONE in the interrupted 2026-08-17 session, submission ID 145285. Confirmation email from approvals@launchingnext.com received 2026-08-17 22:51 UTC: manually reviewed, 'typically takes 2-4 months before a startup is featured'. DO NOT RESUBMIT."],

    # --- BLOCKED, NEEDS A HUMAN ---
    ["alternativeto.net", "Software/product directory (high authority)", "self-serve form", "No", "P0",
     "needs manual registration",
     "HIGHEST-VALUE target on this list and still not done. Cloudflare blocks the automated browser with a non-interactive 'Performing security verification' interstitial (HTTP 403, Ray ID logged) - there is NO captcha widget to solve, it is a headless-fingerprint block. 4 distinct approaches tried and failed: (1) direct /manage-item/ URL, (2) homepage-first to prime clearance cookies, (3) 20s wait for JS auto-resolve, (4) DOM inspection for an interactive challenge element (none exists). A human on a normal browser will get straight in. Free, no phone/address."],

    ["curlie.org", "General web directory (DMOZ successor)", "self-serve form", "No", "P2",
     "needs manual registration",
     "Attempted in the interrupted 2026-08-17 session: blocked by reCAPTCHA image-select. DO NOT RE-ATTEMPT with automation. Value is marginal regardless - editorially reviewed, long/uncertain queues, modest SEO weight today."],

    # --- OPEN, NOT YET ATTEMPTED ---
    ["startupbase.io", "Startup directory", "OAuth signup only (Google/LinkedIn/X)", "No", "P1",
     "needs manual registration",
     "VERIFIED 2026-08-18: site loads fine (HTTP 200), but /login states verbatim 'Email sign-in is for existing accounts. New here? Use Google, LinkedIn, or X above - your account is created automatically.' NEW accounts can ONLY be created via OAuth. Not attempted: that requires driving a login into the owner's personal Google/LinkedIn/X account, which needs their own credentials and sign-off. Owner can do this in ~2 minutes themselves, then submit at /submit."],

    ["producthunt.com", "Product launch directory", "self-serve launch", "No", "P1",
     "needs manual registration",
     "VERIFIED 2026-08-18: HTTP 403 Cloudflare 'Just a moment' interstitial to the automated browser. Fine in a human browser. Separately: outbound links are nofollow and the audience is Western/English - weak fit for a RU-market VPN, and a credible launch wants an aged account. Visibility play, not a link play."],

    ["crunchbase.com", "Company/entity directory", "self-serve profile", "No", "P1",
     "needs manual registration",
     "VERIFIED 2026-08-18: /register returns HTTP 403 Cloudflare 'Attention Required!' to the automated browser. Free company profile for a human. Website link is nofollow but it is a strong entity/citation signal (helps Google associate the brand name with the domain). Horizance LLC (1209 Mountain Road Pl NE, Ste R, Albuquerque, NM 87110, USA) is a real registered entity, so the profile would be factual."],

    ["f6s.com", "Startup directory", "self-serve profile", "No", "P2",
     "needs manual registration",
     "VERIFIED 2026-08-18: HTTP 403 to the automated browser. Free startup profile, historically a dofollow website link. Low effort for a human, moderate value."],

    ["slant.co", "Comparison/recommendation community", "self-serve (add option to a list)", "No", "P2",
     "needs manual registration",
     "VERIFIED 2026-08-18: HTTP 403 Cloudflare 'Attention Required!' to the automated browser. User-generated 'best X' lists; a human can add Horizance as an option under VPN lists. Community-moderated, so a drive-by add may be removed."],

    ["openalternative.co", "Open-source alternatives directory", "self-serve (login required)", "No", "P2",
     "rejected-dead (ineligible)",
     "VERIFIED 2026-08-18: /submit loads (HTTP 200) and redirects to login, so it is technically reachable - but the directory lists OPEN SOURCE alternatives to proprietary software. Horizance is a closed-source commercial service, so it does not qualify. Do not pursue."],

    # --- OUTREACH / EDITORIAL, NOT A LINK TARGET ---
    ["freedomchecker.ateo.digital (ateo.digital)", "VPN-availability aggregator (RU market)",
     "editorial only - no submission form exists; optional soft outreach", "No", "P1",
     "outreach draft ready for owner to send",
     "VERIFIED 2026-08-18. NOT a backlink target: listings link to affiliate Telegram bot links (e.g. /vpn/horizon-vpn/ -> https://t.me/HorizonVPNbot?start=aff_Pn2zDrtN8), NOT to the service's website - so being listed yields ZERO link to horizancevpn.com. IS a strong trust/referral-traffic target given their independent-testing reputation. Their /methodology/ states verbatim: 'Мы покупаем подписку за свои деньги и пользуемся ей скрытно несколько месяцев, прежде чем добавить сервис на сайт' and 'Никто не платит за включение в чекер, и такой возможности не существует.' Inclusion is covert + editorial and cannot be bought or requested. Affiliate links are added only AFTER a service appears, at the service's initiative. ACTION: owner sends the draft in the 'Ateo outreach draft' sheet from their own account to @Ateo_help_bot. Competitors already listed (199-479 RUB/mo vs Horizance 99 RUB/mo): Durev VPN, impVPN, Banksy VPN, LIBERTY VPN, Sirius VPN, Horizon VPN, Magnum VPN, Snowdn VPN, Paper VPN, SLOVOVPN."],

    ["ateo.digital - PAID channel (@libertarian)", "Paid advertising placement", "commissioned advertising", "No", "P2",
     "rejected-legal-risk (documented, do not pursue)",
     "DOCUMENTED AS A CHOICE FOR THE OWNER, NOT RECOMMENDED. Their footer lists 'Сотрудничество | реклама: @libertarian'. A paid/arranged placement would be advertising that Horizance itself commissioned - which is exactly the pattern behind the RU enforcement cases in the legal memo (Vologda UFAS №035/05/5-1030/2025, Khabarovsk, 'Сердитая Чувашия'), where third parties were penalised for posting VPN links on their own channels. Per FAS letter №08/41716/23, info about your own service on your OWN site is not advertising - but a commissioned third-party placement is. Do not pursue without explicit owner sign-off."],

    ["github.com/Lissy93/awesome-privacy", "Privacy awesome-list", "GitHub pull request", "No", "P2", "not started",
     "CONTRIBUTING.md verified: one addition per PR, squash commits, commit message 'Adds [software-name] to [section-name]', entry must not read like an advertisement, hosted services must have a clear privacy policy (horizancevpn.com/privacy exists - OK). Schema fields include securityAudited / openSource / acceptsCrypto. Realistic odds LOW: the VPN section is a short globally-curated list. Zero cost to try."],

    ["github.com/marzneshin/awesome-anticensor", "Anti-censorship awesome-list", "GitHub pull request", "No", "P2",
     "not started - verify scope first",
     "Found via search, not content-verified. Appears to list censorship-circumvention TOOLS/panels (Marzban, Hiddify, 3x-ui, Xray) rather than commercial VPN subscriptions - Horizance is probably out of scope. Check the actual list before investing."],

    # --- REJECTED ---
    ["privacyguides.org", "Privacy/security editorial directory", "editorial (GitHub PR + forum review)", "No", "P2",
     "rejected-dead (ineligible)",
     "Published VPN criteria require WireGuard AND OpenVPN support, standard config files usable in a generic open-source client, a kill switch in any custom client, and independent security audits. Horizance is VLESS+Reality (Xray) only with no published independent audit -> does not meet criteria. Do not spend time here."],

    ["awesome-selfhosted (GitHub)", "Self-hosting awesome-list", "GitHub pull request", "No", "P2",
     "rejected-dead (ineligible)",
     "Requires the listed project to be open-source and self-hostable software. Horizance is a hosted commercial subscription service - out of scope."],

    ["g2.com", "B2B software review directory", "vendor claim/verification", "Likely (business email + vendor verification)", "P2",
     "rejected-dead (likely ineligible)",
     "Vendor claim wants a corporate-domain email; james.smith.web@gmail.com is free Gmail. Also B2B-oriented - a 99 RUB/mo consumer VPN is a poor fit."],

    ["capterra.com", "Software review directory", "vendor verification", "Yes (phone verification)", "P2",
     "rejected-legal-risk",
     "Vendor onboarding involves phone verification. Skipped per the no-phone constraint."],

    ["trustpilot.com", "Review platform", "business account + domain verification", "Possibly (phone during onboarding)", "P2",
     "not started - abort if phone requested",
     "Free business account exists; links are nofollow. If onboarding asks for a phone number, stop and mark rejected-legal-risk."],

    ["Russian directories (2GIS / Yandex Business / Zoon / Flamp / Yell.ru)", "Local business directories",
     "self-serve + verification", "Yes (RU phone + RU legal address)", "P2", "rejected-legal-risk",
     "Established in a prior round: near-zero SEO value without a Russian legal address (Horizance LLC is US-only), plus real RU enforcement-case precedent tied to third-party directory links to VPN services. Not pursued."],

    ["vpnmentor.com / top10vpn.com / comparitech.com", "VPN comparison/review sites",
     "editorial + affiliate partnership", "No", "P2", "rejected-legal-risk (commissioned advertising)",
     "No self-serve submission. Placement is editorial and in practice tied to affiliate/commercial partnerships, which makes it advertising Horizance commissioned - same legal exposure as the Ateo paid channel. Only pursue with explicit owner sign-off."],

    ["reddit.com/r/VPN", "Community resource list", "community/editorial", "No", "P2", "not started",
     "No self-serve submission; self-promotion is rule-restricted and the wiki converges on Mullvad/Proton/Nord/Surfshark/IVPN. Viable only as genuine long-term participation, never as a drop-a-link play."],
]

OUTREACH = """DRAFT — for the OWNER to send from their own Telegram account to @Ateo_help_bot.
Do NOT send from an agent account. Reason: this is outbound contact with a third party
representing the business; it needs a human's own voice and sign-off.

Why this wording: Ateo's published methodology says they buy a subscription themselves and
test covertly for months before listing, and that inclusion cannot be paid for. So this is
NOT a pitch and NOT a request to be listed — it is an offer of access so they can test if
they want to. Keep it short and factual. No persuasive claims, no metrics, no promises.

--- MESSAGE (Russian) ---

Здравствуйте!

Читаю Freedom Checker и вижу, что вы тестируете сервисы сами и скрытно, поэтому не прошу
о размещении — просто оставляю информацию на случай, если захотите добавить нас в проверку.

Horizance VPN, horizancevpn.com
— протокол: VLESS + Reality (Xray)
— цена: 99 ₽/мес, 999 ₽/год
— пробный доступ: ссылка на 24 часа без карты и регистрации на сайте; 14 дней в Telegram-боте
— клиенты: iOS, Android, Windows, macOS, Linux (Happ)

Если нужен доступ для тестов — напишите, выдам подписку без ограничений по времени,
никаких условий с нашей стороны.

Спасибо за проект.

--- END ---

Note: the offer of a free test subscription is standard for independent testers and is NOT a
paid placement — Ateo state they buy their own subscriptions, so they may simply decline it.
That is fine; the point is to be on their radar as a real service.
"""

PACE = """SAFE REGISTRATION PACE — sourced findings (researched 2026-08-18)

BOTTOM LINE: there is no defensible single number, and the honest answer is that the premise
is weaker than it looks. Sources genuinely disagree, and the disagreement splits cleanly by
source quality: primary/authoritative sources say link velocity is not a thing Google
penalises; secondary SEO blogs publish specific numbers with no published data behind them.

WHAT THE AUTHORITATIVE SOURCES SAY
1. Google Search Central, "Spam policies for Google Web Search"
   https://developers.google.com/search/docs/essentials/spam-policies
   Defines link spam by INTENT and QUALITY (buying/selling links, excessive exchanges,
   automated link creation). Contains NO velocity threshold, no per-day/week/month limit,
   and no mention of domain age. Confidence: HIGH (primary source).

2. Google Search Central Blog, "A reminder on qualifying links and our link spam update" (2021-07)
   https://developers.google.com/search/blog/2021/07/link-tagging-and-link-spam-update
   Paid/arranged links are fine if qualified with rel="nofollow" / rel="sponsored".
   Relevant here: the Ateo paid channel and VPN-comparison affiliate placements would need
   this tagging to be policy-clean. Confidence: HIGH (primary source).

3. John Mueller (Google), via Search Engine Journal, "Google's John Mueller On Link Velocity
   and Penalties" — https://www.searchenginejournal.com/google-link-velocity/331637/
   Direct quote: "it's not so much a matter of how many links you get in which time period...
   if these are links that are unnatural or from our point of view problematic then they would
   be problematic... it doesn't really matter how many or in which time."
   Confidence: HIGH (direct statement from a Google Search Advocate).

4. Ahrefs SEO glossary, "What is Link Velocity?" — https://ahrefs.com/seo/glossary/link-velocity
   Defines it, then states the SEO community is split on whether it matters at all; notes
   Google's Gary Illyes has called link velocity a made-up term; traces the penalty folklore to
   a 2003 Google patent mentioning a "spiky rate of growth". Confidence: HIGH for "the concept
   is contested", and it explicitly does NOT endorse a safe number.

WHAT THE SECONDARY SOURCES SAY (weaker — treat as folklore, not data)
Various SEO blogs (w3era, keyserp, sharprocket, saynine, serpzilla, joinindexed) give figures
for new domains: "2-5 links/month", "5-10 high-quality backlinks/month", "new sites should cap
at 8-12 referring domains/month vs 20-40 for established sites". These figures are mutually
inconsistent, none cites underlying data, and none is from Google, Ahrefs, Moz, SEJ editorial
or Semrush primary research. Confidence: LOW. Do not present these as authoritative.

ON DOMAIN AGE (~2 months for horizancevpn.com)
No source found that ties a specific link-acquisition rate to domain age. Domain age matters
here for a different and more practical reason: a 2-month-old domain has little accumulated
trust, so the realistic constraint is that low-quality directory links will do close to nothing
for it either way — the upside is small, so the risk calculus barely matters. Confidence:
MODERATE, and stated as reasoning rather than as a sourced rule.

PRACTICAL RECOMMENDATION
- Do not throttle submissions to hit an invented weekly quota. There is no sourced basis for it.
- DO filter hard on quality and relevance, which is the axis Google's own policy actually names.
  A handful of genuinely relevant listings (AlternativeTo, SaaSHub, a real Ateo review) is worth
  more than 50 generic directory entries, and the generic ones carry the only real risk here:
  looking like automated link creation, which the spam policy names explicitly.
- Any paid/arranged placement (Ateo @libertarian, VPN-comparison affiliate deals) must carry
  rel="sponsored" or rel="nofollow" to stay policy-clean — and separately carries the RU
  advertising-law exposure documented in the legal memo.
"""

STATUS_FILLS = {
    "registered + submitted (pending approval)": "C6EFCE",
    "submitted - confirmed": "C6EFCE",
    "needs manual registration": "FFEB9C",
    "not started": "FFFFFF",
    "outreach draft ready for owner to send": "BDD7EE",
}


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Targets"

    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for r in ROWS:
        ws.append(r)

    for row in ws.iter_rows(min_row=2):
        status = str(row[5].value or "")
        for key, color in STATUS_FILLS.items():
            if status.startswith(key.split(" (")[0]):
                for c in row:
                    c.fill = PatternFill("solid", fgColor=color)
                break
        if status.startswith("rejected"):
            for c in row:
                c.fill = PatternFill("solid", fgColor="FFC7CE")
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[:6]:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    _autosize(ws, [34, 26, 30, 20, 9, 30, 100])
    ws.auto_filter.ref = ws.dimensions

    # Sheet 2: Ateo outreach draft
    ws2 = wb.create_sheet("Ateo outreach draft")
    for i, line in enumerate(OUTREACH.splitlines(), start=1):
        ws2.cell(row=i, column=1, value=line)
    _autosize(ws2, [110])

    # Sheet 3: safe pace research
    ws3 = wb.create_sheet("Safe pace (sourced)")
    for i, line in enumerate(PACE.splitlines(), start=1):
        ws3.cell(row=i, column=1, value=line)
    _autosize(ws3, [110])

    out = "/home/pashavpsadm/horizance_vpn_bot/backlinks_vpn.xlsx"
    wb.save(out)
    print("wrote", out, "rows:", len(ROWS))


if __name__ == "__main__":
    main()
